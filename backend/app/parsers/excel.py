"""Excel parsing helpers with header detection and multi-sheet support."""

from __future__ import annotations

import io
import re
from typing import Any

from app.core.errors import ParserError
from app.core.logging import get_logger
from app.core.text import clean_text

log = get_logger(__name__)


def _normalize_header(value: Any) -> str:
    text = clean_text(str(value)) if value is not None else ""
    text = (text or "").casefold()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def read_workbook(data: bytes, *, filename: str = "") -> dict[str, list[list[Any]]]:
    """Read an ``.xlsx``/``.xls`` file into ``{sheet_name: rows}``.

    Both formats are supported; ``.xls`` uses xlrd, ``.xlsx`` uses openpyxl.
    """
    lower = filename.lower()
    if lower.endswith(".xls") and not lower.endswith(".xlsx"):
        return _read_xls(data)
    try:
        return _read_xlsx(data)
    except Exception as exc:
        log.debug("openpyxl failed (%s); retrying with xlrd", exc)
        return _read_xls(data)


def _read_xlsx(data: bytes) -> dict[str, list[list[Any]]]:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    result: dict[str, list[list[Any]]] = {}
    for name in workbook.sheetnames:
        sheet = workbook[name]
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        result[name] = rows
    workbook.close()
    return result


def _read_xls(data: bytes) -> dict[str, list[list[Any]]]:
    import xlrd

    workbook = xlrd.open_workbook(file_contents=data)
    result: dict[str, list[list[Any]]] = {}
    for name in workbook.sheet_names():
        sheet = workbook.sheet_by_name(name)
        rows: list[list[Any]] = []
        for r in range(sheet.nrows):
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                # Convert Excel date serials to datetimes for downstream parsing.
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        row.append(xlrd.xldate.xldate_as_datetime(cell.value, workbook.datemode))
                        continue
                    except Exception:
                        pass
                row.append(cell.value)
            rows.append(row)
        result[name] = rows
    return result


def find_header_row(
    rows: list[list[Any]], required: list[str], *, max_scan: int = 25
) -> int | None:
    """Locate the header row by looking for the largest overlap with ``required``.

    Both sides are normalized (``U.S. Application No.`` -> ``u s application no``)
    so callers can pass human-readable column names.
    """
    normalized_required = [_normalize_header(r) for r in required]
    best_index: int | None = None
    best_score = 0
    for index, row in enumerate(rows[:max_scan]):
        headers = {_normalize_header(c) for c in row if c is not None}
        score = sum(
            1
            for req in normalized_required
            if req and any(req in header or header in req for header in headers)
        )
        if score > best_score:
            best_score = score
            best_index = index
    # Require at least two matched columns to avoid latching onto a title row.
    threshold = min(2, len(normalized_required)) if normalized_required else 1
    return best_index if best_score >= threshold else None


def rows_to_dicts(
    rows: list[list[Any]],
    *,
    required: list[str] | None = None,
    header_row: int | None = None,
) -> list[dict[str, Any]]:
    """Convert sheet rows into dicts keyed by normalized header names."""
    if not rows:
        return []
    if header_row is None:
        header_row = find_header_row(rows, required or [])
    if header_row is None:
        raise ParserError("Could not locate a header row in the worksheet")

    headers = [_normalize_header(c) for c in rows[header_row]]
    records: list[dict[str, Any]] = []
    for row in rows[header_row + 1 :]:
        if row is None:
            continue
        values = list(row)
        if all(clean_text(str(v)) is None for v in values):
            continue
        record: dict[str, Any] = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            record[header] = values[index] if index < len(values) else None
        # Skip repeated header rows that appear mid-sheet after page breaks.
        if _normalize_header(record.get(headers[0], "")) == headers[0] and headers[0]:
            continue
        records.append(record)
    return records


def pick(record: dict[str, Any], *candidates: str, default: Any = None) -> Any:
    """Fetch the first matching key, allowing partial (substring) header matches.

    Matching is tolerant of the ways a source may garble a header: case, runs of
    whitespace, and words that got joined. The FDA Excel export, for instance,
    emits ``DATE OF FIRST APPLICANTAPPROVAL`` where the PDF says
    ``DATE OF FIRST APPLICANT APPROVAL``.
    """
    for candidate in candidates:
        key = _normalize_header(candidate)
        if key in record:
            value = record[key]
            if clean_text(str(value)) is not None:
                return value
    for candidate in candidates:
        key = _normalize_header(candidate)
        squashed = key.replace(" ", "")
        for header, value in record.items():
            if key and (key in header or squashed in header.replace(" ", "")):
                if clean_text(str(value)) is not None:
                    return value
    return default


def sheet_rows(data: bytes, *, filename: str = "", sheet: str | None = None) -> list[list[Any]]:
    workbook = read_workbook(data, filename=filename)
    if sheet and sheet in workbook:
        return workbook[sheet]
    # Prefer the sheet with the most rows; skip obviously empty sheets.
    best = max(workbook.items(), key=lambda kv: len(kv[1]), default=("", []))
    return best[1]
